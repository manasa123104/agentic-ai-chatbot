from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

doc = Document()

# Narrow margins a bit for readability
for section in doc.sections:
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def heading(text, level=1):
    doc.add_heading(text, level=level)

def para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.size = Pt(11)

def code_block(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(8)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Agentic AI Chatbot")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Project File Guide & Unique Example Prompts")
sr.font.size = Pt(14)
sr.font.color.rgb = RGBColor(0x5F, 0x63, 0x68)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
mr = meta.add_run(
    "GitHub: https://github.com/manasa123104/agentic-ai-chatbot\n"
    "Branch: main · Local UI: http://localhost:8080"
)
mr.font.size = Pt(10)

# ---------- Overview ----------
heading("1. Project overview", 1)
para(
    "This Java Spring Boot project is a multi-agent chatbot. You type a user story; "
    "three bots generate a website, unit tests, and functional tests. You can preview "
    "the site, edit it from chat (colors, fonts, images), and browse previous results "
    "in the left sidebar history."
)
para("How to run the web UI:", bold=True)
code_block(".\\scripts\\run-web.ps1\n# then open http://localhost:8080")

# ---------- Folders ----------
heading("2. Folders — what each contains", 1)

folders = [
    ("(project root)", "Maven project config, README, ignore rules, Excel backlog, and entry scripts."),
    ("scripts/", "PowerShell helpers to start the chatbot, run the pipeline, and run generated tests. Also a small Python helper for the Excel backlog."),
    ("src/main/java/com/agentic/chatbot/", "All application Java source code."),
    ("src/main/java/.../agent/", "The three agent bots and the pipeline orchestrator."),
    ("src/main/java/.../config/", "Startup CLI runner and web MVC configuration."),
    ("src/main/java/.../generator/", "Code that turns a story into HTML pages, controllers, and tests; theme packs for restaurant/travel/etc."),
    ("src/main/java/.../llm/", "Optional OpenAI-compatible LLM client (offline templates work without a key)."),
    ("src/main/java/.../model/", "Request/response DTOs, site style, chat history entry, bot logs."),
    ("src/main/java/.../service/", "Chat history persistence, site style storage, chat-based site editing."),
    ("src/main/java/.../web/", "HTTP controllers: chat UI, pipeline API, history API, site preview, editor, Google demo auth."),
    ("src/main/resources/", "application.yml config, Thymeleaf HTML templates, static CSS."),
    ("src/main/resources/templates/", "Chatbot page, site editor page, Google demo page."),
    ("src/main/resources/static/css/", "Styles for chatbot and editor."),
    ("src/test/java/...", "Unit tests for story analysis."),
    ("generated-app/", "OUTPUT folder created by the bots (controllers, HTML pages, tests). Not hand-written source; regenerated each run."),
    ("data/", "Runtime chat history JSON (gitignored). Survives restarts so Previous results stay available."),
]

for name, desc in folders:
    para(name, bold=True)
    para(desc)

# ---------- Root files ----------
heading("3. Root files", 1)
root_files = [
    ("pom.xml", "Maven build file: Spring Boot 3.4.3, Java 17, Thymeleaf, Jackson, Selenium (test)."),
    ("README.md", "Quick start for web UI, CLI pipeline, and tests."),
    (".gitignore", "Ignores target/, .env, generated-app contents, data/, logs, IDE files."),
    (".env.example", "Sample environment variables: OPENAI_API_KEY, Google OAuth, LLM settings."),
    ("Future_Enhancements_Todo_List.xlsx", "Backlog of future features (40 items) with priorities."),
]

for name, desc in root_files:
    bullet(f"{name} — {desc}")

# ---------- Scripts ----------
heading("4. scripts/", 1)
scripts = [
    ("run-web.ps1", "Starts the chatbot on port 8080 (sets JAVA_HOME / Maven path if needed)."),
    ("run-pipeline.ps1", "Runs the 3-bot pipeline from the command line with -UserStory."),
    ("run-generated-tests.ps1", "Runs Maven tests inside generated-app/."),
    ("create_enhancements_xlsx.py", "Builds the Future Enhancements Excel file."),
]
for name, desc in scripts:
    bullet(f"{name} — {desc}")

# ---------- Java packages ----------
heading("5. Java source files (by package)", 1)

heading("5.1 Application entry", 2)
bullet("AgenticChatbotApplication.java — Spring Boot main() entry point.")

heading("5.2 agent/ — the three bots", 2)
agent_files = [
    ("AgentBot.java", "Interface every bot implements."),
    ("AgentContext.java", "Shared input: user story, output folder, API key."),
    ("AgentResult.java", "Shared output: success, message, files written."),
    ("PipelineOrchestrator.java", "Runs Bot 1 → Bot 2 → Bot 3 and builds the API response."),
    ("UiGeneratorBot.java", "Bot 1 — generates webpages/controllers."),
    ("UnitTestBot.java", "Bot 2 — generates JUnit unit tests."),
    ("FunctionalTestBot.java", "Bot 3 — generates functional/end-to-end style tests."),
]
for name, desc in agent_files:
    bullet(f"{name} — {desc}")

heading("5.3 config/", 2)
bullet("PipelineCommandLineRunner.java — Optional CLI flags (--story, --exit) to run without the web UI.")
bullet("WebConfig.java — Web/MVC setup (static uploads, etc.).")

heading("5.4 generator/", 2)
gen_files = [
    ("StoryAnalyzer.java", "Reads the story and decides screen type/domain (login, landing, booking, etc.)."),
    ("ScreenCodeGenerator.java", "Main writer: single page or multi-page big site into generated-app/."),
    ("BigSiteTemplates.java", "HTML for multi-page sites (home, about, services, pricing, blog, careers, contact) with brand name + About (not the raw user prompt)."),
    ("CompanyWebTemplates.java", "Google/company-style single-page templates (login, landing, forms, dashboard)."),
    ("SiteThemePack.java", "Picks restaurant/travel/hospital/fitness/bank/tech images, colors, and fonts from the story or edit instruction."),
    ("UnitTestGenerator.java", "Writes MockMvc/JUnit unit tests."),
    ("FunctionalTestGenerator.java", "Writes functional test classes."),
]
for name, desc in gen_files:
    bullet(f"{name} — {desc}")

heading("5.5 llm/", 2)
bullet("LlmClient.java — Interface for LLM completion.")
bullet("OpenAiLlmClient.java — Calls OpenAI-compatible API when OPENAI_API_KEY is set; otherwise bots use offline templates.")

heading("5.6 model/", 2)
model_files = [
    ("PipelineRequest.java", "JSON in: userStory, runUnitTests, runFunctionalTests."),
    ("PipelineResponse.java", "JSON out: success, message, previewUrl, historyId, bot steps."),
    ("BotStepLog.java", "One bot’s log (name, success, summary, files, duration)."),
    ("SiteStyle.java", "Colors, fonts, footer, hero/logo image URLs for the site editor."),
    ("ChatHistoryEntry.java", "One saved query + result for the sidebar history."),
]
for name, desc in model_files:
    bullet(f"{name} — {desc}")

heading("5.7 service/", 2)
bullet("ChatHistoryService.java — Saves/loads history to data/chat-history.json (up to 100 entries).")
bullet("SiteStyleService.java — Loads/saves SITE_STYLE.json and uploaded images.")
bullet("SiteEditService.java — Applies chat edit instructions (pink color, restaurant images, etc.) using the previous site as context.")

heading("5.8 web/ — controllers", 2)
web_files = [
    ("ChatController.java", "Serves / chatbot page; also supports form POST /run."),
    ("PipelineApiController.java", "POST /api/pipeline — chat UI calls this; routes to generate OR edit."),
    ("ChatHistoryApiController.java", "GET/DELETE /api/history — list, open, clear previous results."),
    ("GeneratedSiteController.java", "Serves /site/{page} preview with live theme CSS."),
    ("SiteEditorController.java", "Visual /edit page for colors, fonts, images."),
    ("GoogleAuthController.java", "Demo Google sign-in flow."),
]
for name, desc in web_files:
    bullet(f"{name} — {desc}")

# ---------- Resources ----------
heading("6. Resources (UI & config)", 1)
res = [
    ("application.yml", "Port 8080, generated-app output dir, history file path, LLM and Google settings."),
    ("templates/index.html", "Chatbot UI + Previous results sidebar + example chips."),
    ("templates/editor.html", "Manual site editor form."),
    ("templates/google-demo.html", "Google-style sign-in demo page."),
    ("static/css/style.css", "Chatbot + history sidebar styles."),
    ("static/css/editor.css", "Site editor styles."),
    ("StoryAnalyzerTest.java", "Tests that story → screen-type detection works."),
]
for name, desc in res:
    bullet(f"{name} — {desc}")

# ---------- Runtime output ----------
heading("7. Runtime output (not source of truth)", 1)
bullet("generated-app/ — Latest generated website (HTML, Java controllers, tests, SITE_STYLE.json, SITE_META.json).")
bullet("data/chat-history.json — Saved chats for the left sidebar (Previous results).")

# ---------- How history works ----------
heading("8. How to browse history", 1)
para(
    "After you generate or edit sites, each query appears under Previous results in the left sidebar. "
    "Click an item to reopen that query and result. Use Edit on an item to open the site editor. "
    "History is stored in data/chat-history.json so it remains after restart."
)
para("Typical flow:", bold=True)
bullet("1. Open http://localhost:8080")
bullet("2. Paste a Simple example below → Send")
bullet("3. Open /site/home to browse the website")
bullet("4. Optionally type an edit (e.g. I want the page to be pink)")
bullet("5. Click Previous results in the left box to revisit any run")

# ---------- Examples ----------
heading("9. Unique example prompts (Simple → Advanced)", 1)
para(
    "These prompts are unique (not the built-in OliveGrove / SkyTrail chips). "
    "Use them in order to fill your history with varied results you can browse later."
)

heading("9.1 Simple — single page / small asks", 2)
simple = [
    ("Login", "As a user, I want a sign-in page so that I can access my account securely."),
    ("Register", "As a new customer, I want a registration page so that I can create an account."),
    ("Contact form", "As a visitor, I want a contact form page so that I can send a message to support."),
    ("Appointment", "As a patient, I want an appointment booking page so that I can schedule a clinic visit."),
]

for i, (label, prompt) in enumerate(simple, 1):
    para(f"S{i}. {label}", bold=True)
    code_block(prompt)

heading("9.2 Medium — multi-page sites with brand name", 2)
medium = [
    ("Pet boarding", "As a pet-care founder, I want a full multi-page website called PawsHarbor Lodge with teal colors and Montserrat font so pet owners can explore boarding services, pricing, careers, and contact us."),
    ("Bookstore café", "As a shop owner, I want a multi-page website called ChapterSteam Books with purple colors and Playfair font so readers can browse events, membership pricing, blog posts, careers, and contact the store."),
    ("Music school", "As a music academy director, I want a multi-page website called Harmonia Notes with gold colors and Inter font so students can explore courses, tuition pricing, teacher blogs, careers, and enroll via contact."),
    ("Eco cleaning", "As a cleaning-business owner, I want a full multi-page website called LeafShine Home with green colors and Lato font so customers can compare services, pricing, tips blog, careers, and request a quote."),
    ("Surf school", "As a surf instructor, I want a multi-page travel-style website called Saltline Surf Co with blue colors and Poppins font so beginners can view lessons, package pricing, coast blog, careers, and book via contact."),
]

for i, (label, prompt) in enumerate(medium, 1):
    para(f"M{i}. {label}", bold=True)
    code_block(prompt)

heading("9.3 Advanced — themed sites + rich branding", 2)
advanced = [
    ("Wedding planner", "As a wedding planner, I want a complete multi-page website called VelvetVow Studios with pink colors and serif font so couples can view packages, pricing, real-wedding stories, careers, and contact us."),
    ("Esports team", "As an esports manager, I want a multi-page fan website called NeonRift Gaming with dark mode and blue accents so fans can see roster services, merch pricing, match blog, careers, and contact the org."),
    ("Planetarium", "As a museum curator, I want a full website called Orion Dome Planetarium with dark theme and yellow accents so visitors can explore exhibits, ticket pricing, science blog, careers, and contact education staff."),
    ("Plant nursery", "As a nursery owner, I want a complete multi-page website called Fernfolk Gardens with green colors and Open Sans font so shoppers can browse plant care services, pricing, garden blog, careers, and contact delivery."),
    ("Electric scooters", "As a mobility startup founder, I want a full company website called VoltRide Urban with green colors and Poppins font so riders can see plans, pricing, safety blog, careers, and contact support."),
]

for i, (label, prompt) in enumerate(advanced, 1):
    para(f"A{i}. {label}", bold=True)
    code_block(prompt)

heading("9.4 Edit instructions (use AFTER generating a site)", 2)
para(
    "These do not create a brand-new site from scratch. They update your previous website "
    "and also appear in Previous results history."
)
edits = [
    "I want the page to be pink",
    "Change the website color to blue and use Poppins font",
    "Put restaurant images in every box",
    "Use travel images in all boxes",
    "Make it dark mode",
    "Edit the website to use orange colors and Montserrat font",
]
for i, e in enumerate(edits, 1):
    para(f"E{i}.", bold=True)
    code_block(e)

heading("9.5 Suggested history-browsing session", 2)
para("Run these in order so your sidebar fills with clear, different entries:")
bullet("1) S1 Login (simple)")
bullet("2) M1 PawsHarbor Lodge (multi-page)")
bullet("3) A1 VelvetVow Studios (pink wedding theme)")
bullet("4) E1 I want the page to be pink (edit)")
bullet("5) E3 Put restaurant images in every box (edit)")
bullet("6) Open Previous results → click each entry to review")

# ---------- Tips ----------
heading("10. Tips", 1)
bullet("Put a brand after called / named (e.g. called PawsHarbor) so the site title is clear.")
bullet("Say multi-page / full website / restaurant / travel to get a big multi-page site.")
bullet("The live website shows brand name + About — not your raw chat prompt.")
bullet("History lives under Previous results; Clear removes all saved chats.")
bullet("GitHub main: https://github.com/manasa123104/agentic-ai-chatbot")

out = Path(r"C:\Users\91984\Projects\agentic-ai-chatbot\Agentic_AI_Chatbot_File_Guide_and_Examples.docx")
doc.save(out)
print(out)
